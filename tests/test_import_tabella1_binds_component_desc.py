from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import duckdb
import openpyxl

from pe3_qc_app.db.ids import new_id
from pe3_qc_app.db.migrations import apply_migrations
from pe3_qc_app.importer.import_service import import_workbook


class TestImportTabella1BindsComponentDesc(unittest.TestCase):
    def test_binds_desc_when_missing_in_sheet(self) -> None:
        tmp = Path(tempfile.gettempdir()) / "pe3_test_tabella1_bind_desc.xlsx"
        if tmp.exists():
            tmp.unlink()

        wb = openpyxl.Workbook()

        ws_mat = wb.active
        ws_mat.title = "materiali"
        ws_mat.append(["codice mat", "desc materiale"])
        ws_mat.append(["MAT001", "Material 001"])

        ws_lotti = wb.create_sheet("LOTTI_RISULTATI ANALISI")
        headers = [
            "Data Richiesta",
            "Codice Materiale",
            "Descrizione Materiale",
            "CODICE ISO",
            "DESCRIZIONE ISO",
            "LOTTO ISO",
            "DATA ANALISI",
            "Durezza",
        ]
        ws_lotti.append(headers)
        ws_lotti.append(["2026-04-01", "MAT001", "Material 001", "ISO1", "", "ISOLOT", "2026-04-02", 55.2])

        from openpyxl.worksheet.table import Table

        tbl = Table(displayName="Tabella1", ref=f"A1:{openpyxl.utils.get_column_letter(len(headers))}2")
        ws_lotti.add_table(tbl)

        wb.save(tmp)
        try:
            wb.close()
        except Exception:
            pass

        dbp = Path(tempfile.gettempdir()) / "pe3_test_tabella1_bind_desc.duckdb"
        if dbp.exists():
            dbp.unlink()
        con = duckdb.connect(str(dbp))
        apply_migrations(con)

        # Seed master components table
        type_id = new_id()
        con.execute("INSERT INTO component_types(type_id, name, slot_key, active) VALUES (?, 'ISOCIANATO', 'ISO', TRUE)", [int(type_id)])
        comp_id = new_id()
        con.execute(
            "INSERT INTO components(component_id, type_id, code, description, unit, active) VALUES (?, ?, 'ISO1', 'Iso desc', 'KG', TRUE)",
            [int(comp_id), int(type_id)],
        )

        import_workbook(con, tmp, lambda *_: None)

        row = con.execute("SELECT component_desc FROM test_component_lots WHERE type_name='ISO'").fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(str(row[0]), "Iso desc")
        con.close()


if __name__ == "__main__":
    unittest.main()

