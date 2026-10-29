from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import duckdb
import openpyxl

from pe3_qc_app.db.migrations import apply_migrations
from pe3_qc_app.importer.import_service import import_workbook


class TestImportTabella1Composition(unittest.TestCase):
    def test_import_component_lots(self) -> None:
        tmp = Path(tempfile.gettempdir()) / "pe3_test_tabella1_comp.xlsx"
        if tmp.exists():
            tmp.unlink()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LOTTI_RISULTATI ANALISI"

        headers = [
            "Data Richiesta",
            "Codice Materiale",
            "Descrizione Materiale",
            "Codice Lastrina",
            "Data Prod Effettiva",
            "CODICE ISO",
            "DESCRIZIONE\nISO",
            "LOTTO\nISO",
            "CODICE POL",
            "DESCRIZIONE\nPOL",
            "LOTTO\nPOL",
            "CODICE CAT",
            "DESCRIZIONE\nCAT",
            "LOTTO\nCAT",
            "CODICE ADD1",
            "DESCRIZIONE\nADD1",
            "LOTTO\nADD1",
            "CODICE ADD2",
            "DESCRIZIONE\nADD2",
            "LOTTO\nADD2",
            "DATA ANALISI",
            "Durezza",
        ]
        ws.append(headers)
        ws.append(
            [
                "2026-04-01",
                "MAT001",
                "Material 001",
                "L001",
                "30.03.2026",
                "ISO1",
                "Iso desc",
                "ISOLOT",
                "POL1",
                "Pol desc",
                "POLLOT",
                "CAT1",
                "Cat desc",
                "CATLOT",
                "A1",
                "Add1 desc",
                "A1LOT",
                "A2",
                "Add2 desc",
                "A2LOT",
                "2026-04-02",
                55.2,
            ]
        )

        from openpyxl.worksheet.table import Table

        tbl = Table(displayName="Tabella1", ref=f"A1:{openpyxl.utils.get_column_letter(len(headers))}2")
        ws.add_table(tbl)

        wb.save(tmp)
        try:
            wb.close()
        except Exception:
            pass

        dbp = Path(tempfile.gettempdir()) / "pe3_test_tabella1_comp.duckdb"
        if dbp.exists():
            dbp.unlink()
        con = duckdb.connect(str(dbp))
        apply_migrations(con)
        res = import_workbook(con, tmp, lambda *_: None)
        self.assertEqual(res.tests, 1)

        prod_date = con.execute("SELECT prod_date FROM tests").fetchone()[0]
        self.assertEqual(prod_date.isoformat(), "2026-03-30")

        rows = con.execute(
            """
            SELECT type_name, component_code, component_desc, lot_code
            FROM test_component_lots
            ORDER BY type_name
            """
        ).fetchall()
        self.assertEqual(len(rows), 5)
        by_t = {r[0]: (r[1], r[2], r[3]) for r in rows}
        self.assertEqual(by_t["ISO"], ("ISO1", "Iso desc", "ISOLOT"))
        self.assertEqual(by_t["POL"], ("POL1", "Pol desc", "POLLOT"))
        self.assertEqual(by_t["CAT"], ("CAT1", "Cat desc", "CATLOT"))
        self.assertEqual(by_t["ADD1"], ("A1", "Add1 desc", "A1LOT"))
        self.assertEqual(by_t["ADD2"], ("A2", "Add2 desc", "A2LOT"))
        con.close()


if __name__ == "__main__":
    unittest.main()
