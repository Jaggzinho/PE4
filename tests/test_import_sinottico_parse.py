from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import duckdb
import openpyxl

from pe3_qc_app.db.migrations import apply_migrations
from pe3_qc_app.importer.import_service import import_workbook


class TestImportSinotticoParse(unittest.TestCase):
    def test_sinottico_items_insert_and_link(self) -> None:
        tmp = Path(tempfile.gettempdir()) / "pe3_test_import.xlsx"
        if tmp.exists():
            tmp.unlink()

        wb = openpyxl.Workbook()
        ws_sin = wb.active
        ws_sin.title = "SINOTTICO"
        ws_sin["A2"] = "LABEL"
        ws_sin["B2"] = "MAC"
        ws_sin["C2"] = "CODICE"
        ws_sin["D2"] = "DESCRIZIONE"
        ws_sin["G2"] = "CODICE MATERIALE"
        ws_sin["A3"] = "X"
        ws_sin["B3"] = "MAC 1"
        ws_sin["C3"] = "MAT001"
        ws_sin["D3"] = "Material 001"
        ws_sin["G3"] = "MAT001"

        ws_mat = wb.create_sheet("materiali")
        ws_mat.append(["codice mat", "desc materiale", "cod lastrina", "desc last"])
        ws_mat.append(["MAT001", "Material 001", "L001", "Las 001"])

        ws_lotti = wb.create_sheet("LOTTI_RISULTATI ANALISI")
        ws_lotti.append(["Codice Materiale"])
        tbl = openpyxl.worksheet.table.Table(displayName="Tabella1", ref="A1:A1")
        ws_lotti.add_table(tbl)

        wb.save(tmp)
        try:
            wb.close()
        except Exception:
            pass

        dbp = Path(tempfile.gettempdir()) / "pe3_test_import.duckdb"
        if dbp.exists():
            dbp.unlink()
        con = duckdb.connect(str(dbp))
        apply_migrations(con)
        res = import_workbook(con, tmp, lambda *_: None)
        self.assertEqual(res.sinottico_items, 1)
        cnt = con.execute("SELECT COUNT(*) FROM sinottico_items").fetchone()[0]
        self.assertEqual(int(cnt), 1)
        mid = con.execute("SELECT material_id FROM sinottico_items WHERE excel_code='MAT001'").fetchone()[0]
        self.assertIsNotNone(mid)
        con.close()


if __name__ == "__main__":
    unittest.main()
